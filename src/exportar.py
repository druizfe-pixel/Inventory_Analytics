import logging
import pandas as pd
from pathlib import Path

from src.constants import COLUMNAS_EXPORTAR_ABC
from src.constants import COLUMNAS_EXPORTAR_EOQ

logger = logging.getLogger(__name__)

from pathlib import Path

from src.constants import COLUMNAS_EXPORTAR_SLOTTING


OUTPUT_DIR = Path("output")


def crear_directorio_salida():
    """
    Crea la carpeta de salida si no existe.
    """

    OUTPUT_DIR.mkdir(exist_ok=True)

    logger.info(
        "Directorio de salida listo: %s",
        OUTPUT_DIR.resolve()
    )

from openpyxl.utils import get_column_letter

def ajustar_columnas(worksheet):
    """
    Ajusta automáticamente el ancho de las columnas
    y congela la fila de encabezados.
    """

    worksheet.freeze_panes = "A2"

    for columna in worksheet.columns:
        longitud = max(
            len(str(celda.value)) if celda.value is not None else 0
            for celda in columna
        )

        letra = get_column_letter(columna[0].column)

        worksheet.column_dimensions[letra].width = min(longitud + 2, 40)

from openpyxl.styles import Font


from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

def formatear_resumen(worksheet):
    """
    Aplica formato profesional a la hoja Resumen.
    """

    azul = PatternFill(
        fill_type="solid",
        start_color="1F4E78",
        end_color="1F4E78"
    )

    borde = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Encabezados
    for celda in worksheet[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )
        celda.fill = azul
        celda.alignment = Alignment(horizontal="center")
        celda.border = borde

    # Celdas de datos
    for fila in worksheet.iter_rows(min_row=2):

        indicador = fila[0].value
        valor = fila[1]

        fila[0].border = borde
        fila[1].border = borde

        fila[0].alignment = Alignment(horizontal="left")
        fila[1].alignment = Alignment(horizontal="right")

        # Cantidades
        if (
            "Productos" in indicador
            or "EOQ máximo" in indicador
        ):
            valor.number_format = '#,##0'

        # Valores monetarios
        elif "Valor" in indicador:
            valor.number_format = '#,##0.00'

        # Promedios
        else:
            valor.number_format = '#,##0.00'


def crear_resumen(df):
    """
    Genera un resumen ejecutivo del análisis.
    """

    resumen = []

    resumen.append({
        "Indicador": "Productos analizados",
        "Valor": len(df)
    })

    resumen.append({
        "Indicador": "Productos excluidos",
        "Valor": (df["abcxyz"] == "N").sum()
    })

    # ==========================
    # KPIs del modelo
    # ==========================

    resumen.append({
        "Indicador": "EOQ promedio",
        "Valor": round(df["eoq"].mean(), 2)
    })

    resumen.append({
        "Indicador": "EOQ máximo",
        "Valor": round(df["eoq"].max(), 2)
    })

    resumen.append({
        "Indicador": "Stock de seguridad promedio",
        "Valor": round(df["stock_seguridad"].mean(), 2)
    })

    resumen.append({
        "Indicador": "Punto de reorden promedio",
        "Valor": round(df["punto_reorden"].mean(), 2)
    })

    resumen.append({
        "Indicador": "Valor total inventario",
        "Valor": round(
            (df["inventario"] * df["vr_unit"]).sum(),
            2
        )
    })

    resumen.append({
        "Indicador": "Valor promedio por SKU",
        "Valor": round(
            (df["inventario"] * df["vr_unit"]).mean(),
            2
        )
    })

    conteo = df["abcxyz"].value_counts().sort_index()

    for categoria, cantidad in conteo.items():

        resumen.append({
            "Indicador": f"Productos {categoria}",
            "Valor": cantidad
        })

    return pd.DataFrame(resumen)


def exportar_resultados(inventario, slotting, plan_reubicacion):
    """
    Exporta los resultados del análisis a un archivo Excel
    con múltiples hojas.
    """

    crear_directorio_salida()

    archivo = OUTPUT_DIR / "inventario_procesado.xlsx"

    with pd.ExcelWriter(archivo, engine="openpyxl") as writer:

        # =====================================================
        # Inventario completo
        # =====================================================
        inventario.to_excel(
            writer,
            sheet_name="Inventario",
            index=False
        )

        # =====================================================
        # ABC XYZ
        # =====================================================
        columnas_abc = [
            c
            for c in COLUMNAS_EXPORTAR_ABC
            if c in inventario.columns
        ]

        inventario[columnas_abc].to_excel(
            writer,
            sheet_name="ABCXYZ",
            index=False
        )

        # =====================================================
        # EOQ
        # =====================================================
        columnas_eoq = [
            c
            for c in COLUMNAS_EXPORTAR_EOQ
            if c in inventario.columns
        ]

        inventario[columnas_eoq].to_excel(
            writer,
            sheet_name="EOQ",
            index=False
        )

        # =====================================================
        # Stock de Seguridad
        # =====================================================
        columnas_stock = [
            c for c in [
                "codigo",
                "lead_time",
                "nivel_servicio",
                "stock_seguridad",
            ]
            if c in inventario.columns
        ]

        inventario[columnas_stock].to_excel(
            writer,
            sheet_name="StockSeguridad",
            index=False
        )

        # =====================================================
        # Punto de Reorden
        # =====================================================
        columnas_pr = [
            c for c in [
                "codigo",
                "demanda_diaria",
                "stock_seguridad",
                "punto_reorden",
            ]
            if c in inventario.columns
        ]

        inventario[columnas_pr].to_excel(
            writer,
            sheet_name="PuntoReorden",
            index=False
        )

        # =====================================================
        # Slotting
        # =====================================================
        columnas_slotting = [
            c
            for c in COLUMNAS_EXPORTAR_SLOTTING
            if c in slotting.columns
        ]

        slotting[columnas_slotting].to_excel(
            writer,
            sheet_name="Slotting",
            index=False
        )


        plan_reubicacion.to_excel(
            writer,
            sheet_name="Plan_Reubicacion",
            index=False
        )


        # =====================================================
        # Resumen Ejecutivo
        # =====================================================
        resumen = crear_resumen(inventario)

        resumen.to_excel(
            writer,
            sheet_name="Resumen",
            index=False
        )

        # =====================================================
        # Formato
        # =====================================================
        for hoja in writer.sheets.values():
            ajustar_columnas(hoja)

        formatear_resumen(writer.sheets["Resumen"])

    logger.info(
        "Resultados exportados correctamente: %s",
        archivo
    )